r"""Convierte un archivo Markdown simple (el README del laboratorio) a Word con python-docx.

Uso:  python herramientas/md_a_docx.py README.md Guia_del_participante_Dia2.docx
Soporta: títulos, párrafos, **negrita**, *cursiva*, `código`, [enlaces](url), listas (con ☐/☑ para tareas),
tablas, bloques de código, citas (>) y <details>/<summary>. Requiere:  pip install python-docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x00, 0x5A, 0x9E)
GRIS = "F2F2F2"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![*\w])\*[^*\s][^*]*\*(?!\w))")


NUEVO_BLOQUE = re.compile(r"^(```|#{1,4} |\||>|</?(details|summary)|---)")
ITEM = re.compile(r"^\s*([-*]|\d+\.) ")


def sombrear(elemento_pr, color: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    elemento_pr.append(sombra)


def linea_horizontal(parrafo) -> None:
    bordes = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    for clave, valor in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "A6A6A6")):
        inferior.set(qn(f"w:{clave}"), valor)
    bordes.append(inferior)
    parrafo._p.get_or_add_pPr().append(bordes)


def agregar_enlace(parrafo, texto: str, url: str) -> None:
    rel = parrafo.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "005A9E")
    subrayado = OxmlElement("w:u")
    subrayado.set(qn("w:val"), "single")
    props.extend([color, subrayado])
    run.append(props)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    enlace.append(run)
    parrafo._p.append(enlace)


def agregar_inline(parrafo, texto: str, negrita: bool = False, tam: float | None = None) -> None:
    """Agrega texto con formato en línea al párrafo."""
    texto = re.sub(r"</?(?:br|summary|details)\s*/?>", "", texto)
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            agregar_inline(parrafo, trozo[2:-2], negrita=True, tam=tam)
            continue
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x9C, 0x2B, 0x2B)
            run.bold = negrita or None
        elif trozo.startswith("[") and "](" in trozo:
            etiqueta, url = trozo[1:-1].split("](", 1)
            agregar_enlace(parrafo, etiqueta, url)
            continue
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            run = parrafo.add_run(trozo[1:-1])
            run.italic = True
        else:
            run = parrafo.add_run(trozo)
            run.bold = negrita or None
        if tam:
            run.font.size = Pt(tam)


def configurar_documento(doc: Document, pie: str) -> None:
    seccion = doc.sections[0]
    seccion.left_margin = seccion.right_margin = Cm(2)
    seccion.top_margin = seccion.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for nivel, tam in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tam)
        estilo.font.color.rgb = AZUL
    doc.styles["Title"].font.color.rgb = AZUL
    parrafo_pie = seccion.footer.paragraphs[0]
    parrafo_pie.text = pie + "  ·  Página "
    parrafo_pie._p.append(_campo_pagina())
    for run in parrafo_pie.runs:
        run.font.size = Pt(8)


def _campo_pagina():
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    return campo


def agregar_codigo(doc: Document, lineas: list[str]) -> None:
    for linea in lineas or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.3)
        sombrear(p._p.get_or_add_pPr(), GRIS)
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def separar_celdas(linea: str) -> list[str]:
    """Divide una fila de tabla por `|`, respetando los `\\|` escapados y los `|` dentro de `código`."""
    linea = linea.strip()
    if linea.startswith("|"):
        linea = linea[1:]
    if linea.endswith("|") and not linea.endswith("\\|"):
        linea = linea[:-1]
    celdas, actual, en_codigo, i = [], "", False, 0
    while i < len(linea):
        c = linea[i]
        if c == "\\" and linea[i + 1:i + 2] == "|":
            actual += "|"
            i += 2
            continue
        if c == "`":
            en_codigo = not en_codigo
        if c == "|" and not en_codigo:
            celdas.append(actual.strip())
            actual = ""
        else:
            actual += c
        i += 1
    celdas.append(actual.strip())
    return celdas


def agregar_tabla(doc: Document, filas: list[list[str]]) -> None:
    columnas = max(len(f) for f in filas)
    tabla = doc.add_table(rows=0, cols=columnas)
    tabla.style = "Table Grid"
    for indice, fila in enumerate(filas):
        celdas = tabla.add_row().cells
        for j in range(columnas):
            parrafo = celdas[j].paragraphs[0]
            parrafo.paragraph_format.space_after = Pt(2)
            agregar_inline(parrafo, fila[j] if j < len(fila) else "", negrita=indice == 0, tam=9)
            if indice == 0:
                sombrear(celdas[j]._tc.get_or_add_tcPr(), "DCE6F1")
    # Ancho proporcional al contenido de cada columna (17 cm útiles)
    pesos = [min(max(len(f[j]) if j < len(f) else 0 for f in filas), 70) + 6 for j in range(columnas)]
    tabla.autofit = False
    for fila in tabla.rows:
        for j, celda in enumerate(fila.cells):
            celda.width = Cm(17 * pesos[j] / sum(pesos))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def agregar_item(doc: Document, linea: str):
    m = re.match(r"^(\s*)([-*]|\d+\.) (.*)$", linea)
    sangria, marcador, texto = m.groups()
    texto = re.sub(r"^\[ \] ", "☐ ", texto)
    texto = re.sub(r"^\[[xX]\] ", "☑ ", texto)
    p = doc.add_paragraph()
    nivel = len(sangria) // 2
    p.paragraph_format.left_indent = Cm(0.9 + 0.6 * nivel)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(("•" if marcador in "-*" else marcador) + " ")  # numeración manual: cada lista reinicia en 1
    agregar_inline(p, texto)
    return p


def convertir(origen: Path, destino: Path) -> None:
    lineas = origen.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configurar_documento(doc, "Microsoft AI Foundry BCI Training Day · Día 2 · Guía del participante")
    i, titulo_listo, ultimo, cita = 0, False, None, None  # `ultimo`: párrafo de texto/lista que aún admite continuación
    while i < len(lineas):
        linea = lineas[i]
        suelta = linea.strip()
        if ultimo is not None and suelta and not NUEVO_BLOQUE.match(suelta) and not ITEM.match(linea):
            agregar_inline(ultimo, " " + suelta)
            i += 1
            continue
        ultimo = None
        if not suelta.startswith(">"):
            cita = None
        if suelta.startswith("```"):
            bloque, i = [], i + 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            agregar_codigo(doc, bloque)
        elif suelta.startswith("|") and i + 1 < len(lineas) and re.match(r"^\|[\s:|-]+\|?$", lineas[i + 1].strip()):
            filas = [separar_celdas(suelta)]
            i += 2
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(separar_celdas(lineas[i]))
                i += 1
            agregar_tabla(doc, filas)
            continue
        elif re.match(r"^#{1,4} ", suelta):
            nivel = len(suelta.split(" ", 1)[0])
            if nivel == 1 and not titulo_listo:
                p, titulo_listo = doc.add_paragraph(style="Title"), True
            else:
                p = doc.add_paragraph(style=f"Heading {nivel}")
            agregar_inline(p, suelta[nivel + 1:])
        elif suelta in ("---", "***"):
            linea_horizontal(doc.add_paragraph())
        elif suelta.startswith(">"):
            if cita is None:
                cita = doc.add_paragraph()
                cita.paragraph_format.left_indent = Cm(0.6)
                sombrear(cita._p.get_or_add_pPr(), "FFF4CE")
            agregar_inline(cita, ("" if not cita.runs else " ") + suelta.lstrip("> ").strip())
        elif re.match(r"^(\s*)([-*]|\d+\.) ", linea):
            ultimo = agregar_item(doc, linea)
        elif suelta.startswith("<summary>"):
            agregar_inline(doc.add_paragraph(), re.sub(r"</?summary>", "", suelta), negrita=True)
        elif suelta and not re.match(r"^</?details>$", suelta):
            ultimo = doc.add_paragraph()
            agregar_inline(ultimo, suelta)
        i += 1
    doc.save(destino)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Uso: python herramientas\\md_a_docx.py README.md Guia_del_participante_Dia2.docx")
    convertir(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"Generado {sys.argv[2]}")
